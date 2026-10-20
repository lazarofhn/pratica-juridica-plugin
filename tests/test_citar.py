"""Testes do citar.py e do build_docx.py (formatacao-entrega) com transcripts sintéticos.

Os transcripts imitam o formato do Claude Code (.jsonl com tool_use/tool_result) e as
saídas do conector de jurisprudência. Números, partes e textos são fictícios.

Rodar: python -m pytest tests/test_citar.py -q   (ou python -m unittest)
"""
from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "formatacao-entrega" / "scripts"
CITAR = SCRIPTS / "citar.py"
BUILD = SCRIPTS / "build_docx.py"

try:
    import docx  # noqa: F401
    TEM_DOCX = True
except ImportError:
    TEM_DOCX = False


def carregar_citar():
    spec = importlib.util.spec_from_file_location("citar_v13", CITAR)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------------------
# Saídas sintéticas do conector (formatos das bases)
# ---------------------------------------------------------------------------
STJ_BUSCA = """## STJ — 3 ementas (dedupe ativo)
**Consulta**: negativa de prestação jurisdicional
---
### Resultado 1 (ID: 1000000000000000001)
**Relevancia**: 0.9000
**Processo**: REsp 2136610 (RECURSO ESPECIAL)
**Registro**: 202400000001
**Relator**: MINISTRO FICTÍCIO
**Orgao**: SEGUNDA TURMA

**Ementa**:
PROCESSUAL CIVIL. EMBARGOS DE DECLARAÇÃO. OMISSÃO. NEGATIVA DE
PRESTAÇÃO JURISDICIONAL CONFIGURADA.
1. Primeiro item do julgado de número longo, com o texto
quebrado em linhas como vem do conector.
2. Segundo item do julgado de número longo.
3. Recurso especial provido.
---
### Resultado 2 (ID: 1000000000000000002)
**Processo**: REsp 213661 (RECURSO ESPECIAL)
**Relator**: MINISTRO FICTÍCIO

**Ementa**:
TRIBUTÁRIO. JULGADO DE NÚMERO CURTO.
1. Item único do julgado de número curto.
---
### Resultado 3 (ID: 1000000000000000003)
**Processo**: AREsp 2136610 (AGRAVO EM RECURSO ESPECIAL)

**Ementa**:
AGRAVO. JULGADO DE OUTRA CLASSE COM O MESMO NÚMERO.
1. Item do agravo.
---"""

STJ_EDCL = """## STJ — 1 ementas (dedupe ativo)
---
### Resultado 1 (ID: 1000000000000000004)
**Processo**: EDcl no REsp 2136610 (EMBARGOS DE DECLARAÇÃO NO RECURSO ESPECIAL)

**Ementa**:
EMBARGOS DE DECLARAÇÃO NO RECURSO ESPECIAL. AUSÊNCIA DE VÍCIO.
1. Embargos rejeitados.
---"""

STJ_TRUNCADA = """## STJ — 1 ementas (dedupe ativo)
---
### Resultado 1 (ID: 2000000000000000001)
**Processo**: REsp 3000001 (RECURSO ESPECIAL)

**Ementa**:
TRIBUTÁRIO. EMENTA LONGA.
1. Primeiro item completo.
2. Segundo item interrompido [… ementa truncada …]
---"""

STJ_DETALHES = """## STJ — Ementa ID 2000000000000000001

**Processo**: 3000001
**Classe**: REsp
**Descricao Classe**: RECURSO ESPECIAL
**Relator**: MINISTRO FICTÍCIO

### Ementa
TRIBUTÁRIO. EMENTA LONGA.
1. Primeiro item completo.
2. Segundo item interrompido, agora inteiro.
3. Terceiro item, que só o detalhes traz."""

STJ_DEFEITO = """## STJ — 1 ementas (dedupe ativo)
---
### Resultado 1 (ID: 3000000000000000001)
**Processo**: REsp 4000001 (RECURSO ESPECIAL)

**Ementa**:
PROCESSUAL CIVIL. DEFEITO DE CODIFICAÇÃO.
1. Item limpo, sem defeito algum.
2. Conforme a Corte, ?a decisão deve ser fundamentada? em todos os pontos.
---"""

TRF5_BUSCA = """## TRF-5 — 1 acórdão(s)
**Consulta**: seguro de cargas
---
### 1. APELAÇÃO CÍVEL — proc 08000000020244058100 (ID `TRF5:G2:PJE:1:2:3`)
**Score**: 0.5000 | **Órgão**: 4ª TURMA | **Julgamento**: 2024-01-01
**Relator**: Desembargador Federal Fictício

> PROCESSO Nº: 0800000-00.2024.4.05.8100 - APELAÇÃO
>  APELANTE: EMPRESA FICTÍCIA LTDA
>  ADVOGADO: Fulano de Tal
>  APELADO: FAZENDA NACIONAL
>  RELATOR(A): Desembargador Federal Fictício - 4ª Turma
>  EMENTA: TRIBUTÁRIO. JULGADO FICTÍCIO DO TRF-5.
>  I - Primeiro fundamento em algarismo romano.
>  II - Segundo fundamento em algarismo romano.
>  III - Apelação desprovida.
---"""

TRF5_DETALHES = """## TRF-5 — Acórdão 08000000020244058100

**Processo**: 08000000020244058100
**Classe**: APELAÇÃO CÍVEL
**Órgão Julgador**: 4ª TURMA

### Ementa
PROCESSO Nº: 0800000-00.2024.4.05.8100 - APELAÇÃO
 APELANTE: EMPRESA FICTÍCIA LTDA
 ADVOGADO: Fulano de Tal
 APELADO: FAZENDA NACIONAL
 RELATOR(A): Desembargador Federal Fictício - 4ª Turma
 EMENTA: TRIBUTÁRIO. JULGADO FICTÍCIO DO TRF-5.
 I - Primeiro fundamento em algarismo romano.
 II - Segundo fundamento em algarismo romano.
 III - Apelação desprovida, com o fecho que só o detalhes traz."""

TRF1_DETALHES = """## TRF-1 — Acórdão 0000001-00.2020.4.01.3400

**Processo**: 0000001-00.2020.4.01.3400
**Classe**: APELAÇÃO CIVEL (AC)
**Órgão Julgador**: DÉCIMA-TERCEIRA TURMA

### Ementa
DIREITO TRIBUTÁRIO. JULGADO FICTÍCIO DO TRF-1. I. CASO EM EXAME 1. Primeiro item em linha única. II. QUESTÃO EM DISCUSSÃO 2. Segundo item, que cita o art. 1.022 do CPC. III. RAZÕES DE DECIDIR 3. Terceiro item. IV. DISPOSITIVO E TESE 4. Recurso desprovido.

### Decisão
A Turma, por unanimidade, negou provimento à apelação."""

RFB_BUSCA = """## Soluções Cosit (RFB) — 1 resultado(s)
---
### 1. Solução de Consulta Cosit nº 999/2019 — ID `555`
**Score**: 0.6 | **Data**: 2019-05-31

> ASSUNTO: CONTRIBUIÇÃO FICTÍCIA
>
>  EMENTA: CRÉDITOS. INSUMOS FICTÍCIOS.
>
>  Texto da solução de consulta fictícia, com a conclusão verificável.
---"""

FALSO_BASH = """### Resultado 1 (ID: 7)
**Processo**: REsp 7777777 (RECURSO ESPECIAL)

**Ementa**:
TEXTO FALSO VINDO DE SAÍDA DE BASH.
---"""

SAIDA_GRANDE = """## STJ — 1 ementas (dedupe ativo)
---
### Resultado 1 (ID: 6000000000000000001)
**Processo**: REsp 6000001 (RECURSO ESPECIAL)

**Ementa**:
EMENTA COMPLETA DA SAÍDA GRANDE.
1. Item da saída grande.
---"""

OUTRO_PROJETO = """## STJ — 1 ementas (dedupe ativo)
---
### Resultado 1 (ID: 5000000000000000001)
**Processo**: REsp 5555555 (RECURSO ESPECIAL)

**Ementa**:
EMENTA DE CONVERSA DE OUTRO CLIENTE.
---"""


def chamada(tid, nome, entrada, texto, ts, embrulhar=True):
    conteudo = json.dumps({"result": texto}, ensure_ascii=False) if embrulhar else texto
    return [
        {"type": "assistant", "timestamp": ts, "message": {"role": "assistant", "content": [
            {"type": "tool_use", "id": tid, "name": nome, "input": entrada}]}},
        {"type": "user", "timestamp": ts, "message": {"role": "user", "content": [
            {"tool_use_id": tid, "type": "tool_result", "content": conteudo}]}},
    ]


def gravar_jsonl(caminho: Path, linhas):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho, "w", encoding="utf-8") as f:
        for obj in linhas:
            f.write(json.dumps(obj, ensure_ascii=False) + "\n")


MCP = "mcp__juris__"
SID = "11111111-2222-3333-4444-555555555555"


class CitarTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        cls.base = base
        cls.raiz = base / "projects"
        proj = cls.raiz / "C--projeto-ficticio"
        grande = proj / SID / "tool-results" / "toolu_grande.txt"
        grande.parent.mkdir(parents=True, exist_ok=True)
        grande.write_text(SAIDA_GRANDE, encoding="utf-8")
        linhas = []
        linhas += chamada("t1", MCP + "stj_buscar_jurisprudencia", {"params": {"consulta": "x"}},
                          STJ_BUSCA, "2026-10-01T10:00:00.000Z")
        linhas += chamada("t2", MCP + "stj_buscar_jurisprudencia", {"params": {"consulta": "y"}},
                          STJ_TRUNCADA, "2026-10-01T10:01:00.000Z")
        linhas += chamada("t3", MCP + "stj_buscar_jurisprudencia", {"params": {"consulta": "z"}},
                          STJ_DEFEITO, "2026-10-01T10:02:00.000Z")
        linhas += chamada("t4", MCP + "trf5_buscar_jurisprudencia", {"params": {"consulta": "w"}},
                          TRF5_BUSCA, "2026-10-01T10:03:00.000Z")
        linhas += chamada("t5", MCP + "trf1_detalhes_acordao", {"params": {"acordao_pk": "00000010020204013400"}},
                          TRF1_DETALHES, "2026-10-01T10:04:00.000Z")
        linhas += chamada("t6", MCP + "rfb_buscar_solucoes", {"params": {"consulta": "v"}},
                          RFB_BUSCA, "2026-10-01T10:05:00.000Z")
        linhas += chamada("t7", "Bash", {"command": "cat fake.md"}, FALSO_BASH,
                          "2026-10-01T10:06:00.000Z", embrulhar=False)
        linhas.append({"type": "assistant", "timestamp": "2026-10-01T10:07:00.000Z",
                       "message": {"role": "assistant", "content": [{"type": "text", "text":
                                   FALSO_BASH.replace("7777777", "8888888")}]}})
        linhas += chamada("toolu_grande", MCP + "stj_buscar_jurisprudencia", {"params": {"consulta": "g"}},
                          "<persisted-output>\nOutput too large (90KB). Full output saved to: "
                          f"{grande}\n\nPreview (first 2KB):\n### Resultado 1 (ID: 6000000000000000001)\n"
                          "**Processo**: REsp 6000001 (RECURSO ESPECIAL)\n\n**Ementa**:\nPRÉVIA CORTADA\n"
                          "</persisted-output>", "2026-10-01T10:08:00.000Z", embrulhar=False)
        linhas += chamada("t9", MCP + "stj_buscar_jurisprudencia", {"params": {"consulta": "e"}},
                          STJ_EDCL, "2026-10-01T10:09:00.000Z")
        gravar_jsonl(proj / f"{SID}.jsonl", linhas)
        # detalhes rodados por subagentes (transcripts em <sessão>/subagents)
        gravar_jsonl(proj / SID / "subagents" / "agent-1.jsonl",
                     chamada("s1", MCP + "stj_detalhes_ementa", {"params": {"ponto_id": 2000000000000000001}},
                             STJ_DETALHES, "2026-10-01T11:00:00.000Z")
                     + chamada("s2", MCP + "trf5_detalhes_acordao", {"params": {"acordao_pk": "TRF5:G2:PJE:1:2:3"}},
                               TRF5_DETALHES, "2026-10-01T11:01:00.000Z"))
        # conversa de outro projeto: não pode ser lida sem --transcript
        cls.outro = cls.raiz / "C--outro-cliente" / "99999999-0000-0000-0000-000000000000.jsonl"
        gravar_jsonl(cls.outro, chamada("o1", MCP + "stj_buscar_jurisprudencia", {"params": {}},
                                        OUTRO_PROJETO, "2026-10-01T09:00:00.000Z"))
        cls.citar = carregar_citar()

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def setUp(self):
        self.env = patch.dict(os.environ, {"CLAUDE_PROJECTS_DIR": str(self.raiz),
                                           "CLAUDE_CODE_SESSION_ID": SID})
        self.env.start()
        self.addCleanup(self.env.stop)

    def resolve(self, processo, **kw):
        return self.citar.resolve(processo, **kw)[0]

    # --- casamento ---------------------------------------------------------
    def test_numero_inteiro_e_nao_prefixo(self):
        self.assertTrue(self.resolve("REsp 213.661").startswith("TRIBUTÁRIO. JULGADO DE NÚMERO CURTO."))
        self.assertTrue(self.resolve("REsp 2.136.610/PR").startswith("PROCESSUAL CIVIL. EMBARGOS"))
        with self.assertRaises(LookupError):
            self.resolve("REsp 21366")

    def test_classe_distingue_mesmo_numero(self):
        self.assertTrue(self.resolve("AREsp 2136610").startswith("AGRAVO. JULGADO DE OUTRA CLASSE"))
        self.assertTrue(self.resolve("EDcl no REsp 2136610").startswith("EMBARGOS DE DECLARAÇÃO NO RECURSO"))
        with self.assertRaises(LookupError) as ctx:
            self.resolve("2136610")       # sem classe: três julgados possíveis
        self.assertIn("id:", str(ctx.exception))
        self.assertTrue(self.resolve(None, id_="1000000000000000003").startswith("AGRAVO."))

    def test_detalhes_vence_busca_truncada(self):
        texto = self.resolve("REsp 3000001")
        self.assertIn("Terceiro item, que só o detalhes traz.", texto)
        self.assertNotIn("truncada", texto)
        self.assertEqual(self.resolve("REsp 3000001", itens="2"), "2. Segundo item interrompido, agora inteiro.")

    def test_busca_truncada_sem_detalhes_falha(self):
        sessao = self.base / "so_busca.jsonl"
        gravar_jsonl(sessao, chamada("x", MCP + "stj_buscar_jurisprudencia", {}, STJ_TRUNCADA, "2026-10-01T10:00:00Z"))
        with self.assertRaises(ValueError) as ctx:
            self.resolve("REsp 3000001", transcript=str(sessao))
        self.assertIn("_detalhes_", str(ctx.exception))
        self.assertEqual(self.resolve("REsp 3000001", transcript=str(sessao), itens="1"),
                         "1. Primeiro item completo.")

    def test_interrogacao_no_lugar_de_aspa(self):
        with self.assertRaises(ValueError) as ctx:
            self.resolve("REsp 4000001")
        self.assertIn("aspa", str(ctx.exception))
        self.assertEqual(self.resolve("REsp 4000001", itens="1"), "1. Item limpo, sem defeito algum.")

    # --- formatos das bases ---------------------------------------------------
    def test_trf5_detalhes_sem_cabecalho_pje(self):
        texto = self.resolve("AC 0800000-00.2024.4.05.8100")
        self.assertTrue(texto.startswith("EMENTA: TRIBUTÁRIO. JULGADO FICTÍCIO DO TRF-5."), texto)
        self.assertNotIn("APELANTE", texto)
        self.assertIn("com o fecho que só o detalhes traz", texto)
        self.assertEqual(self.resolve("0800000-00.2024.4.05.8100", itens="2"),
                         "II - Segundo fundamento em algarismo romano.")

    def test_trf1_itens_na_mesma_linha_e_fim_da_ementa(self):
        texto = self.resolve("AC 0000001-00.2020.4.01.3400")
        self.assertTrue(texto.endswith("4. Recurso desprovido."), texto)
        self.assertNotIn("A Turma, por unanimidade", texto)
        self.assertEqual(self.resolve("AC 0000001-00.2020.4.01.3400", itens="2"),
                         "2. Segundo item, que cita o art. 1.022 do CPC.")
        self.assertEqual(self.resolve("AC 0000001-00.2020.4.01.3400", itens="3-4"),
                         "3. Terceiro item. 4. Recurso desprovido.")

    def test_rfb_citacao_em_bloco(self):
        texto = self.resolve("Solução de Consulta Cosit nº 999/2019",
                             trecho='de "Texto da solução" ate "verificável."')
        self.assertEqual(texto, "Texto da solução de consulta fictícia, com a conclusão verificável.")

    # --- de onde se lê ---------------------------------------------------------
    def test_ignora_bash_e_texto_do_modelo(self):
        for proc in ("REsp 7777777", "REsp 8888888"):
            with self.assertRaises(LookupError):
                self.resolve(proc)

    def test_saida_grande_lida_do_arquivo_salvo(self):
        self.assertEqual(self.resolve("REsp 6000001"), "EMENTA COMPLETA DA SAÍDA GRANDE. 1. Item da saída grande.")

    def test_outro_projeto_so_com_transcript_explicito(self):
        with self.assertRaises(LookupError):
            self.resolve("REsp 5555555")
        self.assertEqual(self.resolve("REsp 5555555", transcript=str(self.outro)),
                         "EMENTA DE CONVERSA DE OUTRO CLIENTE.")

    def test_cowork_sem_id_de_sessao_usa_o_projeto_do_mount(self):
        mnt = self.base / "cowork" / "mnt"
        gravar_jsonl(mnt / ".claude" / "projects" / "-sessions-abc-mnt-outputs" / "s.jsonl",
                     chamada("c1", MCP + "stj_buscar_jurisprudencia", {}, OUTRO_PROJETO.replace(
                         "5555555", "7070707"), "2026-10-01T09:00:00Z"))
        (mnt / "outputs").mkdir(parents=True, exist_ok=True)
        env = {k: v for k, v in os.environ.items()
               if k not in ("CLAUDE_CODE_SESSION_ID", "CLAUDE_SESSION_ID", "CLAUDE_PROJECTS_DIR")}
        env.update({"USERPROFILE": str(self.base / "home"), "HOME": str(self.base / "home")})
        r = subprocess.run([sys.executable, str(CITAR), "--processo", "REsp 7070707"],
                           capture_output=True, env=env, cwd=mnt / "outputs")
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))
        self.assertEqual(r.stdout.decode("utf-8").strip(), "EMENTA DE CONVERSA DE OUTRO CLIENTE.")

    def test_cache_no_workspace_serve_a_outra_sessao(self):
        ws = self.base / "ws_cache"
        texto = self.resolve("REsp 2136610", workspace=str(ws))
        cache = ws / "citacoes" / "resp-2136610.json"
        self.assertTrue(cache.exists())
        with patch.dict(os.environ, {"CLAUDE_CODE_SESSION_ID": "sessao-sem-nada"}):
            texto2, prov = self.citar.resolve("REsp 2.136.610", workspace=str(ws), detalhes=True)
        self.assertEqual(texto, texto2)
        self.assertTrue(prov["via"].startswith("cache"))

    # --- CLI ---------------------------------------------------------------
    def test_cli_imprime_so_o_texto(self):
        env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
        r = subprocess.run([sys.executable, str(CITAR), "--processo", "REsp 2136610", "--itens", "2",
                            "--grifo", "número longo"], capture_output=True, env=env)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))
        self.assertEqual(r.stdout.decode("utf-8").strip(), "2. Segundo item do julgado de número longo.")
        self.assertIn("busca", r.stderr.decode("utf-8"))
        r = subprocess.run([sys.executable, str(CITAR), "--processo", "REsp 2136610", "--json"],
                           capture_output=True, env=env)
        dados = json.loads(r.stdout.decode("utf-8"))
        self.assertEqual(dados["via"], "busca")
        self.assertIn("1000000000000000001", dados["ids"])
        r = subprocess.run([sys.executable, str(CITAR), "--processo", "REsp 4000001"], capture_output=True, env=env)
        self.assertEqual(r.returncode, 2)
        self.assertIn("ERRO", r.stderr.decode("utf-8"))

    # --- fonte: autos ------------------------------------------------------------
    def test_autos_por_doc_e_por_arquivo(self):
        ws = self.base / "ws_autos"
        ws.mkdir(exist_ok=True)
        (ws / "doc_028_ementa-ementa.md").write_text(
            "---\nseq: 28\ntipo: Ementa\n---\n\n<!-- pág. 12 -->\n\nPROCESSUAL CIVIL. EMENTA DOS AUTOS.\n\n"
            "1. Primeiro item dos autos.\n\n2. Segundo item dos autos.\n", encoding="utf-8")
        texto, _ = self.citar.resolve_autos(str(ws), doc="28", itens="2")
        self.assertEqual(texto, "2. Segundo item dos autos.")
        texto, _ = self.citar.resolve_autos(str(ws), arquivo="doc_028_ementa-ementa.md",
                                            trecho='de "EMENTA DOS" ate "AUTOS."')
        self.assertEqual(texto, "EMENTA DOS AUTOS.")
        with self.assertRaises(LookupError):
            self.citar.resolve_autos(str(ws), arquivo="../fora.md")
        with self.assertRaises(LookupError):
            self.citar.resolve_autos(None, doc="28")

    def test_ancora_repetida(self):
        span = self.citar._extract_span
        repetido = "A tese é esta. Requer a perícia. Outra frase. Requer a perícia. Fim."
        self.assertEqual(span(repetido, "Requer a", "perícia.", "x"), "Requer a perícia.")
        diferente = "Requer a perícia contábil. Outra frase. Requer a perícia médica."
        with self.assertRaises(LookupError) as ctx:
            span(diferente, "Requer a perícia", ".", "x")
        self.assertIn("AMBÍGUA", str(ctx.exception))
        with self.assertRaises(LookupError):
            span(diferente, "Requer a", "inexistente", "x")


@unittest.skipUnless(TEM_DOCX, "python-docx não instalado")
class BuildDocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        CitarTests.setUpClass.__func__(cls)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def build(self, entrada: Path, *extra):
        env = {**os.environ, "CLAUDE_PROJECTS_DIR": str(self.raiz), "CLAUDE_CODE_SESSION_ID": SID,
               "PYTHONIOENCODING": "cp1252"}
        saida = entrada.with_suffix(".docx")
        r = subprocess.run([sys.executable, str(BUILD), str(entrada), "--out", str(saida), *map(str, extra)],
                           capture_output=True, env=env)
        return r, saida

    def test_docx_com_citacao_ref_e_padrao_da_casa(self):
        from docx import Document
        md = self.base / "peca.md"
        md.write_text("@titulo\nRECURSO ESPECIAL\n\n@corpo\nParágrafo do corpo.\n\n@citacao-ref\nfonte: STJ\n"
                      "processo: REsp 2.136.610\nitens: 2\ngrifo: número longo\nref: STJ, REsp 2.136.610\n\n"
                      "@data\nRecife, 1º de outubro de 2026.\n", encoding="utf-8")
        r, saida = self.build(md)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))
        self.assertIn("REsp 2.136.610 → busca", r.stdout.decode("utf-8"))
        doc = Document(saida)
        cit = next(p for p in doc.paragraphs if "Segundo item" in p.text)
        self.assertEqual(cit.text, "“2. Segundo item do julgado de número longo.” (STJ, REsp 2.136.610)")
        self.assertAlmostEqual(cit.paragraph_format.left_indent.cm, 2.0, places=2)
        self.assertTrue(any(run.font.highlight_color for run in cit.runs))
        self.assertEqual(cit.runs[0].font.size.pt, 11)
        corpo = next(p for p in doc.paragraphs if p.text == "Parágrafo do corpo.")
        self.assertEqual(corpo.runs[0].font.size.pt, 12)
        data = next(p for p in doc.paragraphs if p.text.startswith("Recife"))
        self.assertEqual(data.alignment, 3)          # justificado (WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_template_traz_timbrado_e_descarta_corpo(self):
        from docx import Document
        from docx.shared import Cm
        modelo = self.base / "timbrado.docx"
        d = Document()
        d.sections[0].header.paragraphs[0].text = "ESCRITÓRIO FICTÍCIO — TIMBRADO"
        d.sections[0].top_margin = Cm(5)
        d.add_paragraph("TEXTO DO MODELO QUE DEVE SUMIR")
        d.save(modelo)
        md = self.base / "com_timbrado.md"
        md.write_text("@titulo\nMANIFESTAÇÃO\n\n@corpo\nTexto.\n", encoding="utf-8")
        r, saida = self.build(md, "--template", modelo)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))
        out = Document(saida)
        self.assertIn("TIMBRADO", out.sections[0].header.paragraphs[0].text)
        self.assertAlmostEqual(out.sections[0].top_margin.cm, 5, places=1)
        textos = [p.text for p in out.paragraphs]
        self.assertNotIn("TEXTO DO MODELO QUE DEVE SUMIR", textos)
        self.assertIn("MANIFESTAÇÃO", textos)

    def test_workspace_para_fonte_autos(self):
        from docx import Document
        ws = self.base / "processo_ws"
        (ws / "redacao").mkdir(parents=True, exist_ok=True)
        (ws / "doc_030_voto-voto.md").write_text("---\nseq: 30\n---\n\n<!-- pág. 3 -->\n\nTrecho do voto "
                                                 "condutor que será citado.\n", encoding="utf-8")
        bloco = "@citacao-ref\nfonte: autos\ndoc: 30\ntrecho: de \"Trecho do voto\" ate \"citado.\"\n"
        md = ws / "redacao" / "G1-merito.md"            # .md abaixo do workspace: detectado
        md.write_text(bloco, encoding="utf-8")
        r, saida = self.build(md)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))
        self.assertIn("“Trecho do voto condutor que será citado.”", [p.text for p in Document(saida).paragraphs])
        fora = self.base / "rascunho_fora.md"            # .md fora do workspace: precisa de --workspace
        fora.write_text(bloco, encoding="utf-8")
        r, _ = self.build(fora)
        self.assertNotEqual(r.returncode, 0)
        r, saida = self.build(fora, "--workspace", ws)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8"))

    def test_erro_de_citacao_nao_grava_docx(self):
        md = self.base / "ruim.md"
        md.write_text("@titulo\nX\n", encoding="utf-8")
        r, saida = self.build(md)
        self.assertEqual(r.returncode, 0)
        antes = saida.read_bytes()
        md.write_text("@titulo\nX\n\n@citacao-ref\nprocesso: REsp 4000001\n", encoding="utf-8")
        r, _ = self.build(md)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("aspa", r.stderr.decode("utf-8"))
        self.assertEqual(saida.read_bytes(), antes)
        md.write_text("@citacao-ref\nprocesso: REsp 2136610\nitens: 2\ngrifo: trecho inexistente\n", encoding="utf-8")
        r, _ = self.build(md)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("grifo não encontrado", r.stderr.decode("utf-8"))


if __name__ == "__main__":
    unittest.main()
