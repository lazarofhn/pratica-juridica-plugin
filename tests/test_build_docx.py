"""Testes do modelo timbrado no build_docx.py (limpeza e --modelo), com documentos sintéticos.

Rodar: python -m pytest tests/test_build_docx.py -q   (ou: python -m unittest tests.test_build_docx)
"""
import importlib.util
import os
from pathlib import Path
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills" / "formatacao-entrega" / "scripts" / "build_docx.py"

_spec = importlib.util.spec_from_file_location("build_docx_teste", SCRIPT)
bd = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(bd)

TIMBRE = "ESCRITÓRIO FICTÍCIO ADVOGADOS"


def documento_do_escritorio(caminho):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = TIMBRE
    doc.add_paragraph("Conteúdo sigiloso do cliente fictício.")
    doc.core_properties.title = "Parecer do cliente fictício"
    doc.core_properties.comments = "revisão interna"
    doc.save(caminho)


class SemVariavel:
    """Tira PRATICA_JURIDICA_DATA do ambiente durante o teste (ou a troca por `valor`)."""

    def __init__(self, valor=None):
        self.valor = valor

    def __enter__(self):
        self.antigo = os.environ.pop("PRATICA_JURIDICA_DATA", None)
        if self.valor is not None:
            os.environ["PRATICA_JURIDICA_DATA"] = self.valor

    def __exit__(self, *exc):
        os.environ.pop("PRATICA_JURIDICA_DATA", None)
        if self.antigo is not None:
            os.environ["PRATICA_JURIDICA_DATA"] = self.antigo


class LimparModelo(unittest.TestCase):
    def test_preparar_modelo_descarta_corpo_e_propriedades(self):
        with tempfile.TemporaryDirectory() as d:
            origem = os.path.join(d, "original.docx")
            destino = os.path.join(d, "modelos", "casa.docx")
            documento_do_escritorio(origem)
            bd.preparar_modelo(origem, destino)
            modelo = Document(destino)
            self.assertEqual([p.text for p in modelo.paragraphs], [])
            self.assertEqual((modelo.core_properties.title, modelo.core_properties.comments), ("", ""))
            self.assertEqual(modelo.sections[0].header.paragraphs[0].text, TIMBRE)

    def test_documento_bruto_como_template_tambem_e_limpo(self):
        with tempfile.TemporaryDirectory() as d:
            origem = os.path.join(d, "original.docx")
            documento_do_escritorio(origem)
            doc = bd.novo_documento(origem)
            self.assertEqual(doc.core_properties.title, "")
            self.assertNotIn("sigiloso", "".join(p.text for p in doc.paragraphs))
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, TIMBRE)


class AcharModelo(unittest.TestCase):
    def test_pela_variavel_de_ambiente(self):
        with tempfile.TemporaryDirectory() as d:
            os.makedirs(os.path.join(d, "modelos"))
            Document().save(os.path.join(d, "modelos", "casa.docx"))
            with SemVariavel(d):
                self.assertEqual(bd.achar_modelo("casa", os.path.join(d, "peca.md")),
                                 os.path.join(d, "modelos", "casa.docx"))

    def test_subindo_da_entrada_ate_dados_juridicos(self):
        with tempfile.TemporaryDirectory() as d:
            modelos = os.path.join(d, "dados-juridicos", "modelos")
            os.makedirs(modelos)
            Document().save(os.path.join(modelos, "casa.docx"))
            entrada = os.path.join(d, "dados-juridicos", "processos", "1234567-89.2024.4.05.8100",
                                   "redacao", "peca.md")
            os.makedirs(os.path.dirname(entrada))
            with SemVariavel():
                self.assertEqual(bd.achar_modelo("casa", entrada), os.path.join(modelos, "casa.docx"))

    def test_modelo_inexistente_lista_os_disponiveis(self):
        with tempfile.TemporaryDirectory() as d:
            os.makedirs(os.path.join(d, "modelos"))
            Document().save(os.path.join(d, "modelos", "casa.docx"))
            with SemVariavel(d), self.assertRaises(SystemExit) as erro:
                bd.achar_modelo("outro", os.path.join(d, "peca.md"))
            self.assertIn("casa", str(erro.exception))


if __name__ == "__main__":
    unittest.main()
